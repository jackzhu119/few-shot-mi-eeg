#!/usr/bin/env python3
"""Build the separate, unapproved paper-2 review revision, never old submissions."""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
from pathlib import Path
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
import fitz

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript/discovery_revision_20261027"
CSL = OUT / "Third_Party/bmt_vancouver.csl"


def run(command):
    result = subprocess.run(command, cwd=ROOT, text=True, capture_output=True)
    if result.returncode:
        raise RuntimeError("Document export failed:\n" + result.stderr[-8000:])
    return result.stderr


def format_word(path):
    document = Document(path)
    for section in document.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        # Observed official Research Article template, not an invented hard rule.
        section.top_margin = section.bottom_margin = Mm(20)
        section.left_margin = section.right_margin = Mm(16.51)
        footer = section.footer.paragraphs[0]
        footer.alignment = 1
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Bibliography"):
        if name in document.styles:
            style = document.styles[name]
            style.font.name, style.font.size = "Times New Roman", Pt(12)
            style.paragraph_format.line_spacing = 1
            style.paragraph_format.space_after = Pt(6)
    for table in document.tables:
        table.autofit = True
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    for span in paragraph.runs:
                        span.font.size = Pt(9)
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                for span in paragraph.runs:
                    span.bold = True
    for paragraph in document.paragraphs:
        if re.match(r"Table [1-3]\.", paragraph.text):
            paragraph.paragraph_format.keep_with_next = True
    document.save(path)
    return {
        "paragraphs": len(document.paragraphs), "tables": len(document.tables),
        "inline_figures": len(document.inline_shapes),
        "extracted_words": len(" ".join(p.text for p in document.paragraphs).split()),
    }


def submission_source(text):
    """Keep Word tables/equations editable and supply the figures separately."""
    captions = re.findall(r"^\*\*Figure \d+\. .*?(?=\n\s*\n|\Z)", text, flags=re.M | re.S)
    numbers = re.findall(r"^\*\*Figure (\d+)\.", text, flags=re.M)
    if not captions or numbers != [str(i) for i in range(1, len(captions) + 1)]:
        raise ValueError("Complete, sequential figure legends required")
    text = re.sub(r"^!\[.*?\]\(figures/.*?\).*?$", "", text, flags=re.M)
    for caption in captions:
        text = text.replace(caption, "")
    return text.rstrip() + "\n\n## Figure legends\n\n" + "\n\n".join(captions) + "\n"


def journal_checks(main_text, reference_keys, bibliography):
    abstract = main_text.split("## Abstract\n", 1)[1].split("**Keywords:**", 1)[0].strip()
    abstract_words = len(re.findall(r"\S+", re.sub(r"[*#]", "", abstract)))
    structure = re.findall(r"\*\*(Objectives|Methods|Results|Conclusions):\*\*", abstract)
    keywords = [v.strip() for v in re.search(r"\*\*Keywords:\*\*\s*([^\n]+)", main_text).group(1).split(";")]
    running = re.search(r"\*\*(?:Short title|Running head):\*\*\s*([^\n]+)", main_text).group(1)
    figures = re.findall(r"^\*\*Figure (\d+)\.", main_text, re.M)
    tables = re.findall(r"^\*\*Table (\d+)\.", main_text, re.M)
    headings = re.findall(r"^(#{1,6})\s+(.+)$", main_text, re.M)
    checks = {
        "journal": "Biomedical Engineering / Biomedizinische Technik", "article_type": "Research Article",
        "abstract_words_including_structure_labels": abstract_words,
        "abstract_limit": 200, "abstract_structure": structure,
        "keywords": keywords, "keywords_count": len(keywords), "running_head": running,
        "running_head_characters": len(running), "figures": len(figures), "tables": len(tables),
        "combined_figures_tables": len(figures) + len(tables), "cited_references_main": len(reference_keys),
        "maximum_heading_depth_including_title": max(len(h[0]) for h in headings),
        "word_layout": "A4, Times New Roman12pt, single spaced; official-template dimensions",
        "author_approved_this_revision": False, "submission_performed": False,
        "AI_disclosure_retained": "ChatGPT/Codex" in main_text,
        "original_sign_flip_and_Holm_results_reported_in_main": all(
            value in main_text for value in ("0.01025", "0.04102", "0.01953", "0.05859")),
    }
    if abstract_words > 200 or structure != ["Objectives", "Methods", "Results", "Conclusions"]:
        raise ValueError("Journal abstract format failed")
    if not 3 <= len(keywords) <= 6 or keywords != sorted(keywords) or any(x != x.lower() for x in keywords):
        raise ValueError("Journal keyword format failed")
    if len(running) > 75 or len(figures) + len(tables) > 8 or len(reference_keys) > 35:
        raise ValueError("Journal title/figure/table/reference limit failed")
    plain = subprocess.check_output(["pandoc", "-f", "markdown-implicit_figures+tex_math_single_backslash",
                                     "--citeproc", "--bibliography", str(bibliography),
                                     "--csl", str(CSL), "-t", "plain"], input=main_text, text=True)
    checks["full_manuscript_words_including_title_abstract_tables_legends_references"] = len(plain.split())
    if len(plain.split()) > 6000:
        raise ValueError("Conservative complete-manuscript word count exceeds6000")
    return checks


def declaration_table(source, target):
    """Editable table using the official field names/order, with draft status."""
    text = source.read_text()
    sections = re.split(r"^## (.+)\n", text, flags=re.M)
    values = {sections[i]: sections[i + 1].strip() for i in range(1, len(sections), 2)}
    order = ("Acknowledgments", "Research ethics", "Informed consent", "Author contributions",
             "Use of large language models, AI and machine learning tools", "Conflict of interest",
             "Research funding", "Data availability", "Software availability")
    document = Document()
    document.add_heading("Ethical and legal declarations — author-review draft", 0)
    document.add_paragraph("Biomedical Engineering / Biomedizinische Technik")
    document.add_paragraph((OUT / "MAIN.md").read_text().splitlines()[0][2:])
    document.add_paragraph("The author has not yet approved this new version for submission. "
                           "The approval sentence in the contributions field remains pending.")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "STATEMENT", "TEXT"
    for name in order:
        row = table.add_row()
        row.cells[0].text = name
        row.cells[1].text = values[name].replace("**", "")
    document.save(target)
    return format_word(target)


def package_review():
    """Package public derived results/code separately from the complete review."""
    analysis = ROOT / "research_logs/paper2_v3_analysis_20261027"
    data_zip = OUT / "Supplementary_Data_and_Code.zip"
    code = sorted(ROOT.glob("scripts/*paper2*v3*20261027.py"))
    code += [ROOT / "scripts/build_paper2_discovery_revision_20261027.py"]
    code += sorted(ROOT.glob("tests/test_paper2_v3*.py"))
    files = sorted(p for p in analysis.rglob("*") if p.is_file() and p.suffix in (".json", ".csv", ".md"))
    response = ROOT / "research_logs/bmt_review_response_20261027"
    files += sorted(p for p in response.rglob("*") if p.is_file() and p.suffix in (".json", ".csv", ".md"))
    external = ROOT / "research_logs/paper2_external_validation_20261027"
    files += sorted(p for p in external.rglob("*") if p.is_file() and p.suffix in (".json", ".csv", ".md"))
    files += sorted((ROOT / "experiments/external_validation_20261027").glob("*.json"))
    code += sorted((ROOT / "experiments/external_validation_20261027").glob("*.py"))
    code += sorted(ROOT.glob("scripts/*paper2_external*20261027.py"))
    code += sorted(ROOT.glob("tests/test_external_validation*.py"))
    code += [ROOT / "experiments/hpc_recalibration_v2/hpc_recalibration" / p
             for p in ("__init__.py", "classical.py", "data.py")]
    files += [ROOT / "docs" / p for p in (
        "paper2_external_validation_protocol_20261027.md",
        "paper2_external_validation_dataset_review_20261027.md")]
    files += [ROOT / "LICENSE", ROOT / "LICENSING.md",
              ROOT / "LICENSES/MIT.txt", ROOT / "LICENSES/CC-BY-4.0.txt"]
    with zipfile.ZipFile(data_zip, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for path in files + code:
            archive.write(path, str(path.relative_to(ROOT)))
    if data_zip.stat().st_size > 10_000_000:
        raise ValueError("Supplementary data exceeds the journal per-file limit")
    review_zip = OUT / "BMT_Author_Review_Package_20261027.zip"
    review_files = sorted(p for p in OUT.rglob("*") if p.is_file()
                          and p.name not in (review_zip.name, "PACKAGE_SHA256.json"))
    with zipfile.ZipFile(review_zip, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for path in review_files:
            archive.write(path, str(path.relative_to(OUT)))
    manifest = {"status": "author_review_not_submitted_or_approved",
                "raw_EEG_or_checkpoint_tensors_included": False,
                "archives": [{"file": p.name, "bytes": p.stat().st_size,
                              "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
                             for p in (data_zip, review_zip)]}
    (OUT / "PACKAGE_SHA256.json").write_text(json.dumps(manifest, indent=2) + "\n")


def build():
    bibliography = OUT / "references.bib"
    bib = bibliography.read_text()
    citation_keys = set(re.findall(r"@\w+\s*\{\s*([^,]+),", bib))
    main = (OUT / "MAIN.md").read_text()
    abstract = main.split("## Abstract\n", 1)[1].split("**Keywords:**", 1)[0].strip()
    keywords = re.search(r"\*\*Keywords:\*\*\s*([^\n]+)", main).group(0)
    (OUT / "ABSTRACT_FOR_REVIEW.md").write_text(
        main.splitlines()[0] + "\n\nBiomedical Engineering / Biomedizinische Technik — author-review draft.\n\n"
        + abstract + "\n\n" + keywords + "\n")
    outputs, warnings = [], {}
    checks = None
    for name, stem in (("MAIN", "BMT_Manuscript_Review"), ("SUPPLEMENT", "Supplementary_Methods")):
        source = OUT / (name + ".md")
        text = source.read_text()
        used = set(re.findall(r"(?<!\w)@([A-Za-z0-9_:-]+)", text))
        missing = used - citation_keys
        if missing:
            raise ValueError("Unresolved citation keys: " + str(sorted(missing)))
        if name == "MAIN":
            checks = journal_checks(text, used, bibliography)
        common = ["pandoc", str(source), "--from=markdown-implicit_figures+tex_math_single_backslash",
                  "--standalone", "--citeproc",
                  "--bibliography", str(bibliography), "--csl", str(CSL),
                  "--resource-path", str(OUT), "--metadata", "lang=en-US"]
        word = OUT / (stem + ".docx")
        warnings[name + "_word"] = run(common + ["-o", str(word)])
        word_info = format_word(word)
        pdf = OUT / (stem + ".pdf")
        pdf_source = OUT / "Sources" / (name + "_PDF.md")
        pdf_source.parent.mkdir(exist_ok=True)
        pdf_source.write_text(re.sub(r"(?m)^(\*\*Table [1-3]\..*)$",
                                     lambda m: "\\Needspace{16\\baselineskip}\n\n" + m.group(1), text))
        layout = OUT / "Sources/bmt_layout.tex"
        layout.write_text("\\usepackage{needspace}\n")
        pdf_command = list(common)
        pdf_command[1] = str(pdf_source)
        warnings[name + "_pdf"] = run(pdf_command + ["--include-in-header", str(layout),
            "--pdf-engine=xelatex", "-V", "mainfont=DejaVu Serif",
            "-V", "sansfont=DejaVu Sans", "-V", "monofont=DejaVu Sans Mono",
            "-V", "geometry:margin=22mm", "-V", "fontsize=11pt",
            "-V", "colorlinks=true", "-V", "urlcolor=blue", "-o", str(pdf)])
        with fitz.open(pdf) as document:
            all_text = "\n".join(page.get_text() for page in document)
            if len(all_text) < 1000 or "[?]" in all_text or "\ufffd" in all_text:
                raise ValueError("Exported PDF has missing text or unresolved citations")
            info = {"pages": len(document), "text_characters": len(all_text),
                    "empty_pages": [i + 1 for i, p in enumerate(document) if not p.get_text().strip()]}
            if info["empty_pages"]:
                raise ValueError("Unexpected empty PDF pages")
        outputs.append({"source": str(source.relative_to(ROOT)), "citation_keys": sorted(used),
                        "word": word_info, "pdf": info})
        if name == "MAIN":
            prepared = OUT / "Sources/BMT_Manuscript_For_Submission.md"
            prepared.parent.mkdir(exist_ok=True)
            prepared.write_text(submission_source(text))
            submission = OUT / "BMT_Manuscript_Submission.docx"
            commands = list(common)
            commands[1] = str(prepared)
            warnings["submission_word"] = run(commands + ["-o", str(submission)])
            submission_info = format_word(submission)
            if submission_info["inline_figures"] != 0 or submission_info["tables"] != 3:
                raise ValueError("Submission Word must have three editable tables and separate figures")
            checks["submission_word"] = submission_info
    for source_name, stem in (("COVER_LETTER_DRAFT", "Cover_Letter_DRAFT"),
                              ("ETHICAL_AND_LEGAL_DECLARATIONS_DRAFT", "Ethical_and_Legal_Declarations_DRAFT")):
        source = OUT / (source_name + ".md")
        word = OUT / (stem + ".docx")
        if source_name.startswith("ETHICAL"):
            checks["separate_declarations_word"] = declaration_table(source, word)
        else:
            warnings[source_name] = run(["pandoc", str(source), "--standalone", "-o", str(word)])
            format_word(word)
    if any("not found" in warning.lower() for warning in warnings.values()):
        raise ValueError("Pandoc reported a missing resource/reference: " + str(warnings))
    assets = []
    for file in sorted(OUT.rglob("*")):
        if file.is_file() and file.suffix in (".md", ".bib", ".pdf", ".docx", ".png", ".svg", ".csv"):
            assets.append({"file": str(file.relative_to(ROOT)), "bytes": file.stat().st_size,
                           "sha256": hashlib.sha256(file.read_bytes()).hexdigest()})
    record = {"revision": "2026-10-27", "status": "new_author_review_draft_not_submission_approval",
              "old_submission_files_modified": False,
              "document_builder_performed_fitting_or_inference": False,
              "external_experiment_evidence_added": True,
              "pandoc": subprocess.check_output(["pandoc", "--version"], text=True).splitlines()[0],
              "documents": outputs, "warnings": warnings, "assets": assets}
    (OUT / "DOCUMENT_BUILD.json").write_text(json.dumps(record, indent=2) + "\n")
    checks["supplementary_files_below_10MB"] = all(
        (OUT / f"Supplementary_Methods.{suffix}").stat().st_size <= 10_000_000
        for suffix in ("pdf", "docx"))
    (OUT / "JOURNAL_FORMAT_CHECKS.json").write_text(json.dumps(checks, indent=2) + "\n")
    package_review()
    print(json.dumps({"documents": outputs, "assets": len(assets)}, indent=2))


if __name__ == "__main__":
    build()
