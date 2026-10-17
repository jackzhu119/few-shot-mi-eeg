#!/usr/bin/env python3
"""Build the innovation revision in PLOS ONE submission and reading formats.

Run with the bundled artifact Python. --render uses the canonical DOCX renderer
and retains page PNGs for inspection; it does not certify visual QA. Existing
scientific sources and the earlier editorial package are never modified.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from xml.sax.saxutils import escape

import docx
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript/plos_one_innovation_20261017"
BUILD = ROOT / "research_logs/kumar_plos_innovation_20261017/document_build"
BIB = ROOT / "manuscript/kumar_plos_one_references_innovation_20261017.bib"
CSL = ROOT / "manuscript/review_package_author_confirmed_20261015/plos.csl"
RENDERER = Path("/root/.codex/skills/builtins/documents/render_docx.py")
SPECS = (
    ("PLOS_ONE_Manuscript_Submission", "submission_draft", True, False, False),
    ("PLOS_ONE_Manuscript_Reading", "submission_draft", False, True, False),
    ("S1_Supplementary_Methods", "supporting_methods", False, False, False),
    ("S2_Reproducibility", "audit_reproducibility", False, False, False),
    ("Cover_Letter", "cover_letter", False, False, True),
)
CAPTION = re.compile(r"^\*\*Fig\s+(\d+)\.\s", re.I)
TABLE_CAPTION = re.compile(r"^\*\*Table\s+(\d+)\.\s", re.I)
IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^\n]+)\)(?:\{[^}]*\})?\s*$")
GUIDELINES = {
    "manuscript": "https://journals.plos.org/plosone/s/submission-guidelines",
    "figures": "https://journals.plos.org/plosone/s/figures",
    "tables": "https://journals.plos.org/plosone/s/tables",
    "AI_disclosure": "https://journals.plos.org/plosone/s/ethical-publishing-practice",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def relative(path: Path) -> str:
    return str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path)


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def read_source(kind: str) -> tuple[Path, str, str]:
    source = ROOT / f"manuscript/kumar_plos_one_{kind}_innovation_20261017.md"
    first, body = source.read_text(encoding="utf-8").split("\n", 1)
    require(first.startswith("# "), f"Missing Markdown title: {source.name}")
    return source, first[2:].strip(), body.strip() + "\n"


def main_source_checks(title: str, body: str) -> dict:
    short_match = re.search(r"^\*\*Short title:\*\*\s*(.+)$", body, re.M)
    require(short_match is not None, "Main manuscript requires a short title")
    short = short_match.group(1).strip()
    abstract_match = re.search(r"^## Abstract\s*\n(.*?)(?=^## )", body, re.M | re.S)
    require(abstract_match is not None, "Abstract must precede Introduction")
    abstract = abstract_match.group(1).strip()
    words = len(abstract.split())
    require(len(title) <= 250, "Full title exceeds 250 characters")
    require(len(short) <= 100, "Short title exceeds 100 characters")
    require(words <= 300, f"Abstract has {words} words, exceeding 300")
    require(not re.search(r"\[@|@\w+|\[[0-9]+\]", abstract), "Abstract contains citations")
    require(not re.search(r"\b(?:funding|funded|funders?)\b", body, re.I),
            "Move funding information to Author_Submission_Details and the submission form")
    require(not re.search(r"\[\^[^\]]+\]", body), "Main manuscript has footnotes")
    headings = re.findall(r"^(#{2,})\s+(.+)$", body, re.M)
    require(all(len(level) <= 4 for level, _ in headings), "More than three heading levels")
    top = [text.strip() for level, text in headings if level == "##"]
    require(top[:2] == ["Abstract", "Introduction"], "Opening section order")
    require("Acknowledgments" in top, "PLOS ending section requires Acknowledgments")
    require(top.index("Acknowledgments") < top.index("References"), "Acknowledgments order")
    require(top[-1] == "Supporting information captions", "Supporting captions must be last")
    tables = [int(x) for x in re.findall(r"^\*\*Table\s+(\d+)\.", body, re.M)]
    require(tables == list(range(1, len(tables) + 1)), "Table caption order is not consecutive")
    without_captions = re.sub(r"^\*\*(?:Table|Fig)\s+\d+\.[^\n]*$", "", body, flags=re.M)
    missing_table_citations = [n for n in tables
                               if not re.search(rf"\bTable\s+{n}\b", without_captions)]
    require(not missing_table_citations,
            f"Tables require explicit prose citations: {missing_table_citations}")
    table_first_mentions = list(dict.fromkeys(
        int(number) for number in re.findall(r"\bTable\s+(\d+)\b", without_captions)
    ))
    require(table_first_mentions == tables,
            f"First table citations must be consecutive, including figure legends: {table_first_mentions}")
    return {"title": title, "title_characters": len(title), "short_title": short,
            "short_title_characters": len(short), "abstract_words": words,
            "heading_levels": max(len(level) - 1 for level, _ in headings),
            "top_level_sections": top, "table_numbers": tables,
            "table_first_citation_order": table_first_mentions,
            "funding_in_main_text": False, "footnotes_in_main_text": False}


def position_main_figures(body: str, illustrated: bool) -> tuple[str, list[dict]]:
    """Move each complete caption, and optional image, after its first prose citation."""
    blocks = re.split(r"\n\s*\n", body.strip())
    figures, removed = {}, set()
    for index, block in enumerate(blocks):
        match = CAPTION.match(block)
        if not match:
            continue
        number = int(match.group(1))
        require(number not in figures, f"Duplicate Fig {number} caption")
        saved = [block]
        removed.add(index)
        image_block = None
        if index > 0 and IMAGE.match(blocks[index - 1]):
            image_block = blocks[index - 1]
            removed.add(index - 1)
        # A standalone bold title is followed by its full legend in these sources.
        if re.fullmatch(r"\*\*[^*]+\*\*", block):
            require(index + 1 < len(blocks), f"Missing Fig {number} legend")
            legend = blocks[index + 1]
            require(not re.match(r"^(?:#|\*\*(?:Fig|Table)\s|!\[|\|)", legend),
                    f"Fig {number} needs its explanatory legend")
            # A leading panel label is prose, not an alphabetic ordered list.
            legend = re.sub(r"^(\([A-Za-z]\))(?=\s)", r"\\\1", legend)
            saved.append(legend)
            removed.add(index + 1)
        require(image_block is not None, f"Missing reading-copy image for Fig {number}")
        image_match = IMAGE.match(image_block)
        image_path = ROOT / "manuscript" / image_match.group(2)
        require(image_path.exists(), f"Figure source does not exist: {image_path}")
        figures[number] = {"caption": saved, "image": image_block,
                           "image_path": relative(image_path)}
    require(figures, "No figure captions found in manuscript")
    require(sorted(figures) == list(range(1, len(figures) + 1)), "Figure caption sequence")
    remaining = [block for i, block in enumerate(blocks) if i not in removed]
    placed, result, placements = set(), [], []
    for block in remaining:
        result.append(block)
        for number in sorted(figures):
            if number in placed or not re.search(rf"\bFig\s+{number}\b", block):
                continue
            figure = figures[number]
            if illustrated:
                result.append(figure["image"])
            result.extend(figure["caption"])
            placed.add(number)
            placements.append({"figure": number, "image_source": figure["image_path"],
                               "caption_after_first_citation": True,
                               "embedded_image": illustrated})
    require(placed == set(figures), f"Missing prose figure citations: {sorted(set(figures) - placed)}")
    require([x["figure"] for x in placements] == sorted(figures), "First figure citation order")
    return "\n\n".join(result) + "\n", placements


def load_base_builder():
    path = ROOT / "scripts/build_kumar_editorial_documents_20261017.py"
    spec = importlib.util.spec_from_file_location("kumar_editorial_document_base", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def prepare_citation_style() -> Path:
    """Retain the PLOS style and render explicit preprint metadata losslessly."""
    tree = etree.parse(str(CSL))
    ns = {"c": "http://purl.org/net/xbiblio/csl"}
    choose = tree.xpath("//c:macro[@name='container-title']/c:choose", namespaces=ns)[0]
    preprint = etree.fromstring(
        '<if xmlns="http://purl.org/net/xbiblio/csl" variable="genre" type="article-journal" match="all">'
        '<group suffix="."><text variable="container-title"/>'
        '<text variable="issue" prefix=":"/>'
        '<text variable="genre" prefix=" [" suffix="]"/></group></if>'
    )
    choose[0].tag = "{http://purl.org/net/xbiblio/csl}else-if"
    choose.insert(0, preprint)
    target = BUILD / "plos_innovation.csl"
    tree.write(str(target), encoding="utf-8", xml_declaration=True, pretty_print=True)
    return target


def format_document(path: Path, *, submission: bool, illustrated: bool, cover: bool) -> None:
    base = load_base_builder()
    base.format_document(path, submission=submission, illustrated=illustrated, cover=cover)
    document = Document(path)
    for section in document.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.85 if cover else 1)
        section.left_margin = section.right_margin = Inches(1)
        columns = base.element(section._sectPr, "w:cols")
        columns.set(qn("w:num"), "1")
    body_size = 12 if submission else (11 if cover else 11.5)
    for style_name in ("Normal", "Body Text", "First Paragraph", "Compact", "Caption", "Bibliography"):
        if style_name in document.styles:
            style = document.styles[style_name]
            base.font(style, body_size)
            if submission:
                style.paragraph_format.line_spacing = 2
                style.paragraph_format.space_after = Pt(0)
    if "Line Number" in document.styles:
        base.font(document.styles["Line Number"], 8)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if style_name in document.styles and submission:
            style = document.styles[style_name]
            base.font(style, 16 if style_name == "Title" else 12)
            style.paragraph_format.line_spacing = 2
    for style_name in ("Hyperlink", "FollowedHyperlink"):
        if style_name in document.styles:
            base.font(document.styles[style_name], body_size)
            document.styles[style_name].font.underline = False
    previous_figure_caption = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if previous_figure_caption:
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            if illustrated:
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    base.font(run, 10.5)
        previous_figure_caption = bool(re.match(r"^Fig\s+\d+\.", text))
        if paragraph.style.name == "Bibliography":
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.tab_stops.clear_all()
            paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(0.25))
        if submission and paragraph.style.name == "Bibliography":
            paragraph.paragraph_format.line_spacing = 2
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                base.font(run, 12)
        if submission and paragraph.style.name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            paragraph.paragraph_format.line_spacing = 2
            for run in paragraph.runs:
                base.font(run, 16 if paragraph.style.name == "Title" else 12)
        if re.match(r"^(?:Fig\s+\d+|Table\s+\d+|S\d+\s+(?:Text|Data|Table|Fig))\.", text):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(8)
            if submission:
                paragraph.paragraph_format.line_spacing = 2
                for run in paragraph.runs:
                    base.font(run, 12)
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.line_spacing = 1
        # Long repository paths may break at slashes rather than force table/page overflow.
        if paragraph.style.name == "Source Code":
            paragraph.paragraph_format.keep_together = False
    # Scale the inherited table geometry from 166 mm to Letter's 6.5-inch text width.
    table_width = round(6.5 * 1440)
    for table in document.tables:
        old_widths = [int(grid.get(qn("w:w"))) for grid in table._tbl.tblGrid.gridCol_lst]
        ratio = table_width / sum(old_widths)
        widths = [round(value * ratio) for value in old_widths]
        header = " ".join(cell.text for cell in table.rows[0].cells).lower()
        if len(widths) == 4 and "left recall" in header and "right recall" in header:
            widths = [round(table_width * proportion) for proportion in (0.30, 0.70/3, 0.70/3, 0.70/3)]
        widths[-1] += table_width - sum(widths)
        for grid, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            grid.set(qn("w:w"), str(width))
        base.element(table._tbl.tblPr, "w:tblW").set(qn("w:w"), str(table_width))
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Inches(widths[i] / 1440)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        base.font(run, 10.5)
    for shape in document.inline_shapes:
        factor = min(1.0, Inches(6.1) / shape.width, Inches(6.6) / shape.height)
        shape.width, shape.height = int(shape.width * factor), int(shape.height * factor)
    # Pandoc produces editable native Office Math, not equation images.
    settings = document.settings.element
    math_properties = settings.find(qn("m:mathPr"))
    if math_properties is None:
        math_properties = OxmlElement("m:mathPr")
        settings.append(math_properties)
    math_font = base.element(math_properties, "m:mathFont")
    math_font.set(qn("m:val"), "Latin Modern Math")
    document.core_properties.title = document.paragraphs[0].text
    document.core_properties.comments = ""
    document.save(path)


def font_environment(font_dir: Path | None) -> tuple[dict, dict]:
    environment = os.environ.copy()
    if font_dir:
        font_dir = font_dir.resolve()
        require(font_dir.is_dir(), f"Font directory does not exist: {font_dir}")
        configuration = BUILD / "fontconfig.xml"
        cache = Path("/tmp/paper2_document_fontcache_20261017")
        cache.mkdir(exist_ok=True)
        configuration.write_text(
            '<?xml version="1.0"?>\n<!DOCTYPE fontconfig SYSTEM "fonts.dtd">\n'
            '<fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
            f'<dir>{escape(str(font_dir))}</dir><cachedir>{escape(str(cache))}</cachedir>'
            '</fontconfig>\n', encoding="utf-8",
        )
        environment["FONTCONFIG_FILE"] = str(configuration)
    fonts = {}
    for requested in ("Times New Roman", "Arial", "Latin Modern Math"):
        result = subprocess.run(["fc-match", "--format=%{family}|%{file}", requested],
                                env=environment, text=True, capture_output=True, check=True)
        family, file_path = result.stdout.strip().split("|", 1)
        fonts[requested] = {"resolved_family": family, "path": file_path,
                            "sha256": sha256(Path(file_path)), "exact_family": family == requested}
    if font_dir:
        require(fonts["Times New Roman"]["exact_family"], "Times New Roman was not found by Fontconfig")
    return environment, fonts


def structural_audit(path: Path, *, submission: bool) -> dict:
    document = Document(path)
    bibliography = [p for p in document.paragraphs if p.style.name == "Bibliography"]
    line_numbers = [{qn_name.rsplit("}", 1)[-1]: value for qn_name, value in node.attrib.items()}
                    for node in document.element.xpath(".//w:lnNumType")]
    checks = {
        "file": relative(path), "sha256": sha256(path),
        "paragraphs": len(document.paragraphs), "tables": len(document.tables),
        "inline_images": len(document.inline_shapes),
        "native_math_objects": len(document.element.xpath(".//m:oMath")),
        "bibliography_entries": len(bibliography),
        "footnote_references": len(document.element.xpath(".//w:footnoteReference")),
        "line_numbers": line_numbers,
        "single_column": all(s._sectPr.find(qn("w:cols")).get(qn("w:num")) == "1"
                             for s in document.sections),
        "body_font": document.styles["Normal"].font.name,
        "body_size_pt": document.styles["Normal"].font.size.pt,
        "body_line_spacing": document.styles["Normal"].paragraph_format.line_spacing,
        "bibliography_double_spaced": all(p.paragraph_format.line_spacing == 2 for p in bibliography),
        "figure_captions": [p.text for p in document.paragraphs if re.match(r"^Fig\s+\d+\.", p.text)],
        "visual_qa": "pending inspection of every rendered page",
    }
    if submission:
        require(checks["inline_images"] == 0, "Submission main text contains embedded figures")
        require(checks["body_size_pt"] == 12 and checks["body_line_spacing"] == 2, "Body typography")
        require(checks["bibliography_double_spaced"], "References must remain double-spaced")
        require(line_numbers and all(n.get("restart") == "continuous" for n in line_numbers),
                "Continuous line numbers missing")
    return checks


def update_author_details(title: str, short: str) -> None:
    prior = ROOT / "manuscript/plos_one_revision_20261017/Author_Submission_Details.md"
    text = prior.read_text(encoding="utf-8")
    text = re.sub(r"(?m)^\| Full title \|.*$", f"| Full title | {title} |", text)
    text = re.sub(r"(?m)^\| Short title \|.*$", f"| Short title | {short} |", text)
    text = text.replace("Current writing revision:", "Current manuscript and analysis revision:")
    text = text.replace("## Scientific approval and current revision", "## Scientific review record and current revision")
    text = re.sub(r"(?m)^- \*\*Current package:\*\*.*$",
                  "- **Current package:** the 17 October 2026 innovation revision in "
                  "`manuscript/plos_one_innovation_20261017/`. The earlier recorded scientific "
                  "approval refers to the identified 16 October content version; this field "
                  "does not extend that record to newly added analyses.", text)
    text += ("\n## Submission preparation\n\n"
             "Enter the financial disclosure in the submission system; it is excluded from "
             "the main manuscript. The corresponding author's ORCID and individual CRediT "
             "roles above are the recorded submission-field values. The source manuscript "
             "contains the account of AI assistance applicable to this revision.\n")
    (OUT / "Author_Submission_Details.md").write_text(text, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only", choices=[spec[0] for spec in SPECS])
    parser.add_argument("--render", action="store_true", help="Render every selected DOCX to PDF and page PNGs")
    parser.add_argument("--font-dir", type=Path, help="Local licensed font directory; fonts are not copied")
    parser.add_argument("--check", action="store_true", help="Validate main source and report; create no document")
    args = parser.parse_args()
    main_source, main_title, main_body = read_source("submission_draft")
    source_checks = main_source_checks(main_title, main_body)
    if args.check:
        print(json.dumps(source_checks, indent=2, ensure_ascii=False))
        return
    require(shutil.which("pandoc") is not None, "Pandoc is required")
    require(BIB.exists() and CSL.exists(), "Bibliography or PLOS CSL missing")
    if args.render:
        require(RENDERER.exists(), "Canonical DOCX renderer is unavailable")
    OUT.mkdir(parents=True, exist_ok=True)
    BUILD.mkdir(parents=True, exist_ok=True)
    environment, fonts = font_environment(args.font_dir)
    citation_style = prepare_citation_style()
    audit_path = OUT / "formatting_audit_20261017.json"
    prior = json.loads(audit_path.read_text()) if args.only and audit_path.exists() else {}
    records = [x for x in prior.get("documents", []) if Path(x["file"]).stem != args.only]
    for name, kind, submission, illustrated, cover in SPECS:
        if args.only and args.only != name:
            continue
        source, title, body = read_source(kind)
        placements = []
        if kind == "submission_draft":
            body, placements = position_main_figures(body, illustrated)
        render_source = BUILD / f"{name}.md"
        render_source.write_text(body, encoding="utf-8")
        target = OUT / f"{name}.docx"
        command = ["pandoc", str(render_source), "--from=markdown-implicit_figures", "--standalone",
                   "--shift-heading-level-by=-1", "--citeproc", f"--bibliography={BIB}", f"--csl={citation_style}",
                   f"--resource-path={ROOT / 'manuscript'}", "--metadata", f"title={title}",
                   "--metadata", "lang=en-US", "--metadata", "link-citations=true", "--output", str(target)]
        result = subprocess.run(command, cwd=ROOT, env=environment, capture_output=True, text=True)
        (BUILD / f"{name}.pandoc.stderr").write_text(result.stderr)
        result.check_returncode()
        require("citation not found" not in result.stderr.lower(), "Unresolved bibliography citation")
        format_document(target, submission=submission, illustrated=illustrated, cover=cover)
        record = structural_audit(target, submission=submission)
        record.update({"source": relative(source), "source_sha256": sha256(source),
                       "prepared_source_sha256": sha256(render_source), "figure_placement": placements})
        if args.render:
            pages = BUILD / "rendered" / name
            command = [sys.executable, str(RENDERER), str(target), "--output_dir", str(pages),
                       "--emit_pdf", "--dpi", "150"]
            rendered = subprocess.run(command, cwd=ROOT, env=environment, capture_output=True, text=True)
            (BUILD / f"{name}.render.stdout").write_text(rendered.stdout)
            (BUILD / f"{name}.render.stderr").write_text(rendered.stderr)
            rendered.check_returncode()
            rendered_pdf = pages / f"{name}.pdf"
            require(rendered_pdf.exists(), "Renderer did not emit PDF")
            pdf = target.with_suffix(".pdf")
            shutil.copyfile(rendered_pdf, pdf)
            import fitz
            with fitz.open(pdf) as rendered_document:
                record["pdf"] = {"file": relative(pdf), "sha256": sha256(pdf),
                                 "pages": len(rendered_document),
                                 "actual_fonts": sorted({f[3] for page in rendered_document for f in page.get_fonts()}),
                                 "page_images": relative(pages)}
                if cover:
                    require(len(rendered_document) == 1, "Cover letter exceeds one page; shorten its source")
        records.append(record)
    update_author_details(main_title, source_checks["short_title"])
    audit = {"schema_version": 1, "checked_against": GUIDELINES,
             "source_checks": source_checks, "main_source_sha256": sha256(main_source),
             "bibliography_sha256": sha256(BIB), "csl_sha256": sha256(CSL),
             "prepared_csl": relative(citation_style), "prepared_csl_sha256": sha256(citation_style),
             "runtime": {"python": sys.version.split()[0], "python_executable": sys.executable,
                         "python_docx": docx.__version__,
                         "pandoc": subprocess.check_output(["pandoc", "--version"], text=True).splitlines()[0]},
             "font_resolution": fonts, "font_binaries_copied_to_release": False,
             "documents": records,
             "status": "structural_checks_complete_visual_review_pending",
             "note": "Page images require inspection before release. No journal submission was performed."}
    audit_path.write_text(json.dumps(audit, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({"built": len(records), "directory": str(OUT), "audit": str(audit_path)}, indent=2))


if __name__ == "__main__":
    main()
