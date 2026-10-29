"""Export the three-author component paper from saved evidence, without fitting."""
from __future__ import annotations

import hashlib
import importlib.util
import json
from pathlib import Path
import re
import subprocess
import zipfile

from docx import Document
from docx.shared import RGBColor
import fitz

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'manuscript/bmt_component_revision_20261029'
BASE = ROOT / 'manuscript/bmt_submission_20261028'
AUDIT = ROOT / 'research_logs/paper2_structure_revision_20261029'
SPEC = importlib.util.spec_from_file_location('document_helpers', ROOT / 'scripts/build_paper2_discovery_revision_20261027.py')
HELPER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(HELPER)


def run(command):
    r = subprocess.run(command, cwd=ROOT, text=True, capture_output=True)
    if r.returncode:
        raise RuntimeError('Document export failed: ' + r.stderr[-6000:])
    return r.stderr


def styled_word(path):
    info = HELPER.format_word(path)
    d = Document(path)
    for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4']:
        if name in d.styles:
            d.styles[name].font.color.rgb = RGBColor(0, 0, 0)
            d.styles[name].font.name = 'Times New Roman'
    for p in d.paragraphs:
        if re.match(r'(?:Table|Figure)\s+(?:S?\d+)\.', p.text):
            p.paragraph_format.keep_with_next = True
    d.core_properties.author = 'Ziyuan Zhu; Sihao Liu; Haoyang Yu'
    d.save(path)
    info['editable_math_objects'] = len(d.element.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
    return info


def declarations():
    text = (OUT / 'Ethical_and_Legal_Declarations.md').read_text()
    parts = re.split(r'^## (.+)\n', text, flags=re.M)
    values = {parts[i]: parts[i + 1].strip() for i in range(1, len(parts), 2)}
    d = Document()
    d.add_heading('Ethical and legal declarations', 0)
    status = json.loads((OUT / 'SUBMISSION_STATUS.json').read_text())
    d.add_paragraph('Biomedical Engineering / Biomedizinische Technik — ' + status['display_version'])
    d.add_paragraph('Ziyuan Zhu; Sihao Liu; Haoyang Yu')
    d.add_paragraph((OUT / 'MAIN.md').read_text().splitlines()[0][2:])
    table = d.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = 'STATEMENT'
    table.rows[0].cells[1].text = 'TEXT'
    for name in ['Acknowledgments', 'Research ethics', 'Informed consent', 'Author contributions',
                 'Use of large language models, AI and machine learning tools', 'Conflict of interest',
                 'Research funding', 'Data availability', 'Software availability']:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = name, values[name].replace('**', '')
    p = OUT / 'Ethical_and_Legal_Declarations.docx'
    d.save(p)
    return styled_word(p)


def package():
    data = OUT / 'Supplementary_Data_and_Code.zip'
    selected = {}
    with zipfile.ZipFile(BASE / 'Supplementary_Data_and_Code.zip') as previous:
        for name in previous.namelist():
            if not name.endswith('/'):
                selected[name] = previous.read(name)
    directories = [AUDIT, ROOT / 'docs/paper2_structure_revision_20261029']
    for directory in directories:
        for p in directory.rglob('*'):
            if p.is_file() and p.suffix in {'.md', '.json', '.csv', '.gz'}:
                selected[str(p.relative_to(ROOT))] = p.read_bytes()
    for p in [ROOT / 'scripts/review_kumar_component_saved_evidence_20261029.py',
              ROOT / 'scripts/render_component_revision_figures_20261029.py',
              ROOT / 'scripts/render_component_external_figure_20261029.py',
              ROOT / 'scripts/render_component_workflow_20261029.py',
              ROOT / 'scripts/build_bmt_component_revision_20261029.py',
              OUT / 'FIGURE_REPRODUCTION.md',
              ROOT / 'research_logs/kumar_cross_session_20261012/protocol/split_manifest.json',
              ROOT / 'LICENSE', ROOT / 'LICENSING.md', ROOT / 'LICENSES/MIT.txt',
              ROOT / 'LICENSES/CC-BY-4.0.txt',
              ROOT / 'manuscript/plos_one_innovation_20261017/Figures/Fig6_participant_source_data.csv',
              ROOT / 'manuscript/plos_one_innovation_20261017/Figures/Fig5_source_data.csv',
              ROOT / 'manuscript/discovery_revision_20261027/figures/Fig1_update_components_source_data.csv']:
        selected[str(p.relative_to(ROOT))] = p.read_bytes()
    for p in (ROOT / 'research_logs/paper2_hpc_v2_20261023/existing_evidence').glob('*'):
        if p.suffix in {'.csv', '.json', '.md'}:
            selected[str(p.relative_to(ROOT))] = p.read_bytes()
    for p in (OUT / 'figures').glob('*.csv'):
        selected[str(p.relative_to(ROOT))] = p.read_bytes()
    for p in (OUT / 'workflow').glob('*'):
        if p.suffix in {'.csv', '.json', '.txt'}:
            selected[str(p.relative_to(ROOT))] = p.read_bytes()
    with zipfile.ZipFile(data, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for name, content in sorted(selected.items()):
            if name.endswith(('.npz', '.pt', '.pth', '.edf', '.bdf', '.mat')):
                raise ValueError('EEG or model arrays cannot enter the public supplement')
            z.writestr(name, content)
    if data.stat().st_size > 10_000_000:
        raise ValueError('Supplementary data/code exceeds the journal 10 MB limit')
    upload = OUT / 'BMT_Upload_Files_20261029.zip'
    entries = [('Main_Manuscript.docx', OUT / 'BMT_Manuscript_Submission.docx'),
               ('Ethical_and_Legal_Declarations.docx', OUT / 'Ethical_and_Legal_Declarations.docx'),
               ('Cover_Letter.docx', OUT / 'Cover_Letter.docx'),
               ('Supplementary_Methods.pdf', OUT / 'Supplementary_Methods.pdf'),
               ('Supplementary_Data_and_Code.zip', data)]
    for i, stem in [(1, 'Fig1_components'), (2, 'Fig2_budgets_sessions'), (3, 'Fig3_external_validation')]:
        entries.append((f'Figure{i}.tiff', OUT / 'figures' / (stem + '.tiff')))
    with zipfile.ZipFile(upload, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for name, path in entries:
            z.write(path, name)
    drawing = OUT / 'Publication_Figures_and_Source_20261029.zip'
    figure_sources = {ROOT / 'scripts' / name for name in [
        'render_component_revision_figures_20261029.py',
        'render_component_external_figure_20261029.py',
        'render_component_workflow_20261029.py']}
    figure_sources.update(ROOT / item['path'] for item in json.loads(
        (OUT / 'figures/FIGURE_MANIFEST_20261029.json').read_text())['source_files'])
    figure_sources.update([OUT / 'MAIN.md', OUT / 'references.bib', ROOT / 'LICENSE',
                           ROOT / 'LICENSING.md', OUT / 'FIGURE_REPRODUCTION.md',
                           ROOT / 'research_logs/paper2_external_validation_20261027/aggregate.json',
                           ROOT / 'research_logs/paper2_external_validation_20261027/run_metrics.json'])
    for directory in (OUT / 'figures', OUT / 'workflow'):
        figure_sources.update(p for p in directory.glob('*') if p.is_file())
    with zipfile.ZipFile(drawing, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for p in sorted(figure_sources):
            z.write(p, str(p.relative_to(ROOT)))
        readme = zipfile.ZipInfo.from_file(OUT / 'FIGURE_REPRODUCTION.md', 'README.md')
        readme.compress_type = zipfile.ZIP_DEFLATED
        z.writestr(readme, (OUT / 'FIGURE_REPRODUCTION.md').read_bytes())
    complete = OUT / 'BMT_Three_Author_Review_Package_20261029.zip'
    files = sorted(p for p in OUT.rglob('*') if p.is_file() and p.name not in {
        complete.name, upload.name, drawing.name, 'PACKAGE_SHA256.json', 'PACKAGE_MANIFEST.json'})
    with zipfile.ZipFile(complete, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for p in files:
            z.write(p, str(p.relative_to(OUT)))
    status = json.loads((OUT / 'SUBMISSION_STATUS.json').read_text())
    manifest = {'status': status['status'],
                'final_new_artifact_approval_recorded': status['final_new_artifact_approval_recorded'],
                'submission_performed': False, 'EEG_or_weights_included': False,
                'archives': [{'file': p.name, 'bytes': p.stat().st_size,
                              'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
                             for p in [data, upload, drawing, complete]]}
    (OUT / 'PACKAGE_SHA256.json').write_text(json.dumps(manifest, indent=2) + '\n')


def build():
    main = (OUT / 'MAIN.md').read_text()
    bib = OUT / 'references.bib'
    keys = set(re.findall(r'@\w+\s*\{\s*([^,]+),', bib.read_text()))
    used = set(re.findall(r'(?<!\w)@([A-Za-z0-9_:-]+)', main))
    checks = HELPER.journal_checks(main, used, bib)
    status = json.loads((OUT / 'SUBMISSION_STATUS.json').read_text())
    checks.update({'review_date': '2026-10-29', 'authors': status['authors'],
                   'third_author_membership_and_COI_user_confirmed': True,
                   'all_authors_final_new_artifact_approval_recorded': status['final_new_artifact_approval_recorded'],
                   'reviewer_independence_for_third_author': status['reviewer_independence_for_third_author'],
                   'no_new_fit_or_forward': True})
    abstract = main.split('## Abstract\n', 1)[1].split('**Keywords:**', 1)[0].strip()
    (OUT / 'ABSTRACT_FOR_SUBMISSION.txt').write_text(re.sub(r'\*\*', '', abstract) + '\n')
    sources = OUT / 'Sources'
    sources.mkdir(exist_ok=True)
    (sources / 'bmt_layout.tex').write_text('\\usepackage{needspace}\n')
    rendered, warnings = {}, {}
    plans = [('MAIN.md', 'BMT_Manuscript_Reading'), ('SUPPLEMENT.md', 'Supplementary_Methods'),
             ('Cover_Letter.md', 'Cover_Letter'), ('Ethical_and_Legal_Declarations.md', 'Ethical_and_Legal_Declarations')]
    for source, stem in plans:
        text = (OUT / source).read_text()
        citations = set(re.findall(r'(?<!\w)@([A-Za-z0-9_:-]+)', text))
        if citations - keys:
            raise ValueError('Unresolved citations: ' + str(citations - keys))
        common = ['pandoc', str(OUT / source), '--from=markdown-implicit_figures+tex_math_single_backslash',
                  '--standalone', '--citeproc', '--bibliography', str(bib),
                  '--csl', str(OUT / 'Third_Party/bmt_vancouver.csl'),
                  '--resource-path', str(OUT), '--metadata', 'lang=en-US']
        word = OUT / (stem + '.docx')
        if source == 'Ethical_and_Legal_Declarations.md':
            word_info = declarations()
        else:
            warnings[stem + '_word'] = run(common + ['-o', str(word)])
            word_info = styled_word(word)
        pdf_source = sources / (stem + '.md')
        # Group table captions with their first rows and figure captions with each full image.
        text = re.sub(r'(?m)^(\*\*Table (?:S?\d+)\..*)$', lambda m:'\\Needspace{12\\baselineskip}\n\n' + m.group(1), text)
        text = re.sub(r'(?m)^(!\[Figure[^\n]+)$', lambda m:'\\Needspace{40\\baselineskip}\n\n' + m.group(1), text)
        # Embed vector artwork in PDF; editable Word retains its PNG previews.
        text = re.sub(r'(!\[[^\]]*\]\()([^\)]+)\.png(\))',
                      lambda m: m.group(1) + m.group(2) + '.pdf' + m.group(3)
                      if (OUT / (m.group(2) + '.pdf')).is_file() else m.group(0), text)
        pdf_source.write_text(text)
        pdf_cmd = list(common)
        pdf_cmd[1] = str(pdf_source)
        warnings[stem + '_pdf'] = run(pdf_cmd + ['--include-in-header', str(sources / 'bmt_layout.tex'),
            '--pdf-engine=xelatex', '-V', 'mainfont=DejaVu Serif', '-V', 'sansfont=DejaVu Sans',
            '-V', 'monofont=DejaVu Sans Mono', '-V', 'geometry:margin=22mm', '-V', 'fontsize=11pt',
            '-V', 'colorlinks=true', '-V', 'urlcolor=blue', '-o', str(OUT / (stem + '.pdf'))])
        with fitz.open(OUT / (stem + '.pdf')) as pdf:
            extracted = '\n'.join(p.get_text() for p in pdf)
            if '[?]' in extracted or '\ufffd' in extracted or any(not p.get_text().strip() for p in pdf):
                raise ValueError('Broken/empty PDF: ' + stem)
            pdf_info = {'pages': len(pdf), 'text_characters': len(extracted)}
        rendered[stem] = {'word': word_info, 'pdf': pdf_info}
    submission = sources / 'BMT_Manuscript_For_Submission.md'
    submission.write_text(HELPER.submission_source(main))
    run(['pandoc', str(submission), '--from=markdown-implicit_figures+tex_math_single_backslash', '--standalone',
         '--citeproc', '--bibliography', str(bib), '--csl', str(OUT / 'Third_Party/bmt_vancouver.csl'),
         '--resource-path', str(OUT), '--metadata', 'lang=en-US', '-o', str(OUT / 'BMT_Manuscript_Submission.docx')])
    rendered['BMT_Manuscript_Submission'] = styled_word(OUT / 'BMT_Manuscript_Submission.docx')
    if rendered['BMT_Manuscript_Submission']['tables'] != 2 or rendered['BMT_Manuscript_Submission']['inline_figures']:
        raise ValueError('Submission Word must have two editable tables and separate figures')
    text_doc = Document(OUT / 'BMT_Manuscript_Submission.docx')
    all_text = ' '.join(p.text for p in text_doc.paragraphs)
    all_text += ' ' + ' '.join(c.text for t in text_doc.tables for r in t.rows for c in r.cells)
    checks['exported_complete_word_count'] = len(all_text.split())
    if checks['exported_complete_word_count'] > 6000:
        raise ValueError('Complete exported manuscript exceeds 6000 words')
    if any('not found' in v.lower() for v in warnings.values()):
        raise ValueError('Missing document resources')
    checks['documents'] = rendered
    checks['warnings'] = warnings
    checks['supplement_under_10MB'] = all((OUT / ('Supplementary_Methods.' + ext)).stat().st_size < 10_000_000 for ext in ['docx','pdf'])
    if not checks['supplement_under_10MB']:
        raise ValueError('Supplement exceeds 10 MB')
    (OUT / 'JOURNAL_FORMAT_CHECKS.json').write_text(json.dumps(checks, indent=2) + '\n')
    package()
    print(json.dumps({'format': checks, 'archives': json.loads((OUT / 'PACKAGE_SHA256.json').read_text())}, indent=2))


if __name__ == '__main__':
    build()
