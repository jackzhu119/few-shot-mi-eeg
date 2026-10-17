"""Build the paper-2 editorial revision from Markdown and versioned assets."""
from __future__ import annotations
import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript/plos_one_revision_20261017"
BUILD = ROOT / "research_logs/kumar_editorial_revision_20261017/document_build"
BIB = ROOT / "manuscript/kumar_plos_one_references_20261017.bib"
CSL = ROOT / "manuscript/review_package_author_confirmed_20261015/plos.csl"

def element(parent, name):
    found = parent.find(qn(name))
    if found is None:
        found = OxmlElement(name)
        parent.append(found)
    return found

def font(target, size):
    target.font.name = "Times New Roman"
    target.font.size = Pt(size)
    target.font.color.rgb = RGBColor(0, 0, 0)
    rpr = target.element.get_or_add_rPr()
    color = element(rpr, "w:color")
    color.set(qn("w:val"), "000000")
    for field in ("themeColor", "themeTint", "themeShade"):
        color.attrib.pop(qn("w:" + field), None)
    rf = element(rpr, "w:rFonts")
    for field in ("ascii", "hAnsi", "cs"):
        rf.set(qn("w:" + field), "Times New Roman")
    rf.set(qn("w:eastAsia"), "Noto Serif CJK SC")
    for field in ("asciiTheme", "hAnsiTheme", "csTheme", "eastAsiaTheme"):
        rf.attrib.pop(qn("w:" + field), None)

def format_document(path, *, submission, illustrated, cover):
    doc = Document(path)
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(22)
        section.left_margin = section.right_margin = Mm(22)
        section.header_distance = section.footer_distance = Mm(10)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        element(footer._p.get_or_add_pPr(), "w:suppressLineNumbers")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        field.append(run)
        footer._p.append(field)
        if submission:
            ln = OxmlElement("w:lnNumType")
            for key, value in {"countBy":"1","start":"1","restart":"continuous","distance":"180"}.items():
                ln.set(qn("w:"+key), value)
            sec = section._sectPr
            after = {qn("w:"+x) for x in ("pgNumType","cols","formProt","vAlign","noEndnote","titlePg","textDirection","bidi","rtlGutter","docGrid","printerSettings")}
            pos = next((i for i, child in enumerate(sec) if child.tag in after), len(sec))
            sec.insert(pos, ln)
    normal_size = 12 if submission else (11 if cover else 11.5)
    for name in ("Normal","Body Text","First Paragraph","Compact","Caption","Bibliography"):
        if name in doc.styles:
            st = doc.styles[name]
            font(st, normal_size)
            st.paragraph_format.line_spacing = 2 if submission else (1.05 if cover else 1.18)
            st.paragraph_format.space_after = Pt(5 if cover else 6)
            st.paragraph_format.widow_control = True
    sizes = {"Title":16 if cover else 18,"Subtitle":12,"Heading 1":14,"Heading 2":12.5,"Heading 3":11.5}
    for name, size in sizes.items():
        if name in doc.styles:
            st = doc.styles[name]
            font(st, size)
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.line_spacing = 1.05
            pr = st.element.get_or_add_pPr()
            for border in list(pr.findall(qn("w:pBdr"))):
                pr.remove(border)
    for name in ("Source Code","Verbatim Char","VerbatimChar"):
        if name in doc.styles:
            st = doc.styles[name]
            st.font.name = "DejaVu Sans Mono"
            st.font.size = Pt(8.5)
            if hasattr(st,"paragraph_format"):
                st.paragraph_format.line_spacing = 1.05
    previous_figure_caption = False
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        txt = p.text.strip()
        if txt.startswith("Table ") or re.match(r"(?:Fig [1-4]|S[12] Text|S1 Data)\.",txt):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
        if previous_figure_caption:
            p.paragraph_format.keep_together = True
            if illustrated:
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    font(run,10.5)
        previous_figure_caption = bool(re.match(r"Fig [1-4]\.",txt))
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(5)
        if submission and txt == "Abstract":
            p.paragraph_format.page_break_before = True
        if p.style.name == "Bibliography":
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:
                run.font.size = Pt(10.5)
        if len(txt)>=40 and re.fullmatch(r"[0-9a-f]{40,64}",txt):
            for run in p.runs:
                run.font.name = "DejaVu Sans Mono"
                run.font.size = Pt(8.5)
        if p.style.name in sizes:
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                font(run,sizes[p.style.name])
        for border in list(p._p.get_or_add_pPr().findall(qn("w:pBdr"))):
            p._p.pPr.remove(border)
    for i, shape in enumerate(doc.inline_shapes):
        ratio = shape.height / shape.width
        width = 140 if illustrated and i==2 else 155
        shape.width = Mm(width)
        shape.height = int(Mm(width)*ratio)
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        n = len(table.columns)
        widths = {2:[68,98],3:[60,53,53],4:[62,22,43,39],5:[46,30,30,30,30],6:[30,16,30,30,30,30]}.get(n,[166/n]*n)
        first = " ".join(c.text for c in table.rows[0].cells).lower()
        if n==4 and "source training" in first and "source-validation" in first: widths=[36,58,35,37]
        if n==4 and "unadjusted" in first: widths=[88,22,28,28]
        if n==4 and "source-query" in first: widths=[55,22,48,41]
        if n==4 and "cohort" in first and "cpu source" in first: widths=[50,18,52,46]
        for grid,width in zip(table._tbl.tblGrid.gridCol_lst,widths):
            grid.set(qn("w:w"),str(round(width/25.4*1440)))
        pr = table._tbl.tblPr
        tw = element(pr,"w:tblW")
        tw.set(qn("w:type"),"dxa")
        tw.set(qn("w:w"),str(round(166/25.4*1440)))
        borders = element(pr,"w:tblBorders")
        for side in ("top","left","bottom","right","insideH","insideV"):
            edge = element(borders,"w:"+side)
            for key,value in {"val":"single","sz":"4","color":"D9D9D9"}.items(): edge.set(qn("w:"+key),value)
        for ri,row in enumerate(table.rows):
            rp = row._tr.get_or_add_trPr()
            element(rp,"w:cantSplit")
            if ri==0: element(rp,"w:tblHeader")
            for ci,cell in enumerate(row.cells):
                cell.width = Mm(widths[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = cell._tc.get_or_add_tcPr()
                margins = element(cp,"w:tcMar")
                for side,value in (("top",70),("bottom",70),("left",65),("right",65)):
                    m = element(margins,"w:"+side)
                    m.set(qn("w:w"),str(value))
                    m.set(qn("w:type"),"dxa")
                if ri==0:
                    element(cp,"w:shd").set(qn("w:fill"),"E9EEF3")
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.keep_with_next = False
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(0,0,0)
                        if ri==0: run.bold=True
    doc.core_properties.author = "Ziyuan Zhu"
    doc.core_properties.subject = "Longitudinal motor-imagery EEG component recalibration"
    doc.core_properties.keywords = "EEG, CSP, recalibration, longitudinal decoding"
    order = {name:i for i,name in enumerate(("begChr","sepChr","endChr","grow","shp","ctrlPr"))}
    for pr in doc.element.xpath(".//m:dPr"):
        pr[:] = sorted(list(pr),key=lambda child:order.get(etree.QName(child).localname,99))
    doc.save(path)

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    BUILD.mkdir(parents=True,exist_ok=True)
    specs = [
        ("PLOS_ONE_Manuscript_Revised","submission_draft",False,True,False),
        ("PLOS_ONE_Manuscript_Submission","submission_draft",True,False,False),
        ("S1_Supplementary_Methods_Revised","supporting_methods",False,False,False),
        ("S2_Reproducibility_Revised","audit_reproducibility",False,False,False),
        ("Cover_Letter_Revised","cover_letter",False,False,True),
    ]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only",choices=[s[0] for s in specs])
    args = parser.parse_args()
    record_path = BUILD/"build_record.json"
    records = json.loads(record_path.read_text()) if args.only and record_path.exists() else []
    records = [r for r in records if Path(r["file"]).stem!=args.only]
    created = 0
    for name,kind,submission,illustrated,cover in specs:
        if args.only and name!=args.only: continue
        source = ROOT/f"manuscript/kumar_plos_one_{kind}_20261017.md"
        title,body = source.read_text().split("\n",1)
        title = title.lstrip("# ").strip()
        if not illustrated:
            body = re.sub(r"^!\[[^\]]*\]\([^\n]+\)(?:\{[^}]*\})?\s*$","",body,flags=re.M)
        render = BUILD/f"{name}.md"
        render.write_text(body)
        target = OUT/f"{name}.docx"
        cmd = ["pandoc",str(render),"--from=markdown-implicit_figures","--standalone","--shift-heading-level-by=-1","--citeproc",f"--bibliography={BIB}",f"--csl={CSL}",f"--resource-path={ROOT/'manuscript'}","--metadata",f"title={title}","--metadata","lang=en-US","--metadata","link-citations=true","--output",str(target)]
        result = subprocess.run(cmd,cwd=ROOT,capture_output=True,text=True)
        (BUILD/f"{name}.stderr").write_text(result.stderr)
        result.check_returncode()
        format_document(target,submission=submission,illustrated=illustrated,cover=cover)
        records.append({"file":str(target.relative_to(ROOT)),"source":str(source.relative_to(ROOT)),"source_sha256":hashlib.sha256(source.read_bytes()).hexdigest(),"docx_sha256":hashlib.sha256(target.read_bytes()).hexdigest(),"illustrated":illustrated,"double_spaced_body":submission,"continuous_line_numbers":submission})
        created+=1
    record_path.write_text(json.dumps(records,indent=2)+"\n")
    print(json.dumps({"created":created,"directory":str(OUT)}))

if __name__=="__main__":
    main()
